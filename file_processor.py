"""
File Processor — Excel/CSV থেকে Uid এক্সট্রাক্ট ও রিপোর্ট পার্স
"""
import openpyxl
import csv


def extract_uids(file_path, file_name):
    """
    ইউজারের ফাইল থেকে শুধু Uid (A কলাম) বের করে
    """
    uids = []
    ext = file_name.lower().split(".")[-1]

    try:
        if ext in ["csv", "txt"]:
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                reader = csv.reader(f)
                for row in reader:
                    if row and row[0].strip():
                        val = row[0].strip()
                        if val.lower() not in ["uid", "id", "username", "user"]:
                            uids.append(val)

        elif ext in ["xlsx", "xls"]:
            wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
            for sheet in wb.worksheets:
                for row in sheet.iter_rows(values_only=True):
                    if row and row[0] is not None:
                        val = str(row[0]).strip()
                        if val and val.lower() not in ["uid", "id", "username", "user"]:
                            uids.append(val)
            wb.close()

    except Exception as e:
        print(f"extract_uids error: {e}")

    # Duplicate remove
    return list(dict.fromkeys(uids))


def parse_report(file_path, file_name):
    """
    অ্যাডমিনের রিপোর্ট পার্স করে
    রিটার্ন: dict { uid: 'ok' or 'of' }
    """
    report = {}
    ext = file_name.lower().split(".")[-1]

    GOOD = ["ok", "good", "active", "1", "yes", "valid", "live"]
    BAD = ["of", "bad", "dead", "0", "no", "off", "invalid", "die"]

    try:
        if ext in ["xlsx", "xls"]:
            wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
            for sheet in wb.worksheets:
                for row in sheet.iter_rows(values_only=True):
                    if len(row) >= 3 and row[0] and row[2]:
                        uid = str(row[0]).strip()
                        status_raw = str(row[2]).strip().lower()

                        if status_raw in GOOD:
                            report[uid] = "ok"
                        elif status_raw in BAD:
                            report[uid] = "of"
            wb.close()

        elif ext in ["csv", "txt"]:
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                reader = csv.reader(f)
                for row in reader:
                    if len(row) >= 3 and row[0] and row[2]:
                        uid = row[0].strip()
                        status_raw = row[2].strip().lower()
                        if status_raw in GOOD:
                            report[uid] = "ok"
                        elif status_raw in BAD:
                            report[uid] = "of"

    except Exception as e:
        print(f"parse_report error: {e}")

    return report
